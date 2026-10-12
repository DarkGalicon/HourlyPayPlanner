from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from hourly_pay.domain.enums import SessionPeriod, StudyForm, WarningCode, WorkType
from hourly_pay.domain.errors import UnrecognizedFormatError
from hourly_pay.importers.nagr_xlsx_v1 import NagrXlsxV1Importer, calculate_sha256


def test_imports_title_decimal_rows_without_double_counting(synthetic_workbook: Path) -> None:
    result = NagrXlsxV1Importer().import_file(synthetic_workbook)

    assert result.teacher.full_name == "Тестов Т.Т."
    assert result.teacher.department == "Кафедра тестирования"
    assert result.academic_year.name == "2026-2027"
    assert result.source.employment_type == "штатный преподаватель"
    assert result.source.reported_rate == Decimal("0.5")
    assert result.source.reported_total_hours == Decimal("8")
    assert result.source.include_in_fact is None
    assert result.normalized_hours == Decimal("8")
    assert len(result.items) == 4
    assert {item.work_type for item in result.items} == {
        WorkType.LECTURE,
        WorkType.CONTROL,
        WorkType.UNCLASSIFIED,
    }
    assert all(item.work_type is not WorkType.STATE_EXAM for item in result.items)


def test_anomalies_are_reported_and_idpo_is_recognized(synthetic_workbook: Path) -> None:
    result = NagrXlsxV1Importer().import_file(synthetic_workbook)
    codes = [warning.code for warning in result.warnings]

    assert WarningCode.UNEXPLAINED_TOTAL_DIFF in codes
    assert WarningCode.UNKNOWN_CONTROL_KIND in codes
    assert WarningCode.NON_EVEN_AUDITORY_HOURS in codes
    assert WarningCode.MISSING_GROUP in codes
    assert WarningCode.MISSING_DISCIPLINE in codes
    assert WarningCode.SUMMARY_TOTAL_MISMATCH not in codes
    assert any(item.is_idpo for item in result.items)
    unexplained = next(
        warning
        for warning in result.warnings
        if warning.code is WarningCode.UNEXPLAINED_TOTAL_DIFF
    )
    assert "Итого 6 ч" in unexplained.message
    assert "не подтверждена" in unexplained.message


def test_known_control_kinds_are_not_anomalies(synthetic_workbook: Path) -> None:
    result = NagrXlsxV1Importer(known_control_kinds=("Эк", "Зч", "Неизвестно")).import_file(
        synthetic_workbook
    )
    assert WarningCode.UNKNOWN_CONTROL_KIND not in {warning.code for warning in result.warnings}


def test_exam_consultation_outside_row_total_is_not_double_counted(
    synthetic_workbook: Path,
) -> None:
    workbook = load_workbook(synthetic_workbook)
    workbook["AutumnOch"]["S5"] = "4"
    workbook["AutumnOch"]["U5"] = "2"
    workbook["Summary"]["B3"] = "6"
    workbook["Summary"]["F3"] = "6"
    workbook.save(synthetic_workbook)

    result = NagrXlsxV1Importer().import_file(synthetic_workbook)

    consultation = next(
        item for item in result.items if item.work_type is WorkType.EXAM_CONSULTATION
    )
    assert consultation.hours == Decimal("2")
    assert consultation.raw_context["outside_row_total"] is True
    assert result.normalized_hours == Decimal("6")
    assert result.idpo_hours == Decimal("4")
    assert WarningCode.OUTSIDE_ROW_TOTAL in {warning.code for warning in result.warnings}
    assert WarningCode.SUMMARY_TOTAL_MISMATCH not in {
        warning.code for warning in result.warnings
    }


def test_numbered_row_with_only_total_is_counted_as_unclassified(
    synthetic_workbook: Path,
) -> None:
    workbook = load_workbook(synthetic_workbook)
    sheet = workbook["AutumnOch"]
    sheet["A7"] = "3"
    sheet["B7"] = "Г-3"
    sheet["G7"] = "Руководство подготовкой"
    sheet["S7"] = "20"
    workbook["Summary"]["B3"] = "28"
    workbook["Summary"]["F3"] = "28"
    workbook.save(synthetic_workbook)

    result = NagrXlsxV1Importer().import_file(synthetic_workbook)

    unclassified = next(
        item
        for item in result.items
        if item.source_row_number == 7 and item.work_type is WorkType.UNCLASSIFIED
    )
    assert unclassified.hours == Decimal("20")
    assert result.normalized_hours == Decimal("28")
    assert any(
        warning.code is WarningCode.UNEXPLAINED_TOTAL_DIFF
        and warning.source_row_number == 7
        for warning in result.warnings
    )


def test_duplicate_sha_is_reported(synthetic_workbook: Path) -> None:
    digest = calculate_sha256(synthetic_workbook)
    result = NagrXlsxV1Importer().import_file(synthetic_workbook, known_hashes={digest})
    assert WarningCode.DUPLICATE_IMPORT in {warning.code for warning in result.warnings}


def test_incompatible_workbook_has_controlled_error(tmp_path: Path) -> None:
    path = tmp_path / "unknown.xlsx"
    Workbook().save(path)
    with pytest.raises(UnrecognizedFormatError, match="Формат файла не распознан"):
        NagrXlsxV1Importer().import_file(path)


def test_private_samples_cover_employment_types_and_correspondence() -> None:
    paths = sorted(Path("tests/fixtures_private").glob("*.xlsx"))
    if not paths:
        pytest.skip("Private XLSX fixtures are not present")

    results = [NagrXlsxV1Importer().import_file(path) for path in paths]
    assert {result.source.employment_type for result in results} == {
        "штатный преподаватель",
        "внутренний совместитель",
        "почасовик",
    }
    assert {result.source.reported_rate for result in results} == {
        Decimal("1"),
        Decimal("0.5"),
        Decimal("0.33"),
    }
    assert {result.academic_year.name for result in results} == {"2026-2027"}
    assert all(
        item.semester == 1
        for result in results
        for item in result.items
        if item.study_form is StudyForm.FULL_TIME and item.raw_context["sheet"] == "AutumnOch"
    )
    assert all(
        item.semester == 2
        for result in results
        for item in result.items
        if item.study_form is StudyForm.FULL_TIME and item.raw_context["sheet"] == "SpringOch"
    )
    correspondence = [
        item
        for result in results
        for item in result.items
        if item.study_form is StudyForm.CORRESPONDENCE
    ]
    assert correspondence
    assert all(item.semester is None for item in correspondence)
    assert {item.session_period for item in correspondence if item.session_period} == {
        SessionPeriod.UST,
        SessionPeriod.WINTER,
        SessionPeriod.SUMMER,
    }
    assert all(
        result.normalized_hours == result.source.reported_total_hours
        for result in results
    )
    assert all(
        warning.code is not WarningCode.SUMMARY_TOTAL_MISMATCH
        for result in results
        for warning in result.warnings
    )
    assert sum(
        (item.hours for result in results for item in result.items
         if item.raw_context.get("outside_row_total")),
        start=Decimal("0"),
    ) == Decimal("22")
    assert sum(
        (item.hours for result in results for item in result.items
         if item.work_type is WorkType.UNCLASSIFIED),
        start=Decimal("0"),
    ) == Decimal("26")
