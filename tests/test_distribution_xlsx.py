from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from hourly_pay.application.monthly_planning import MonthlyPlanningRow
from hourly_pay.domain.enums import StudyForm, WorkType
from hourly_pay.domain.errors import InvalidDistributionExportError
from hourly_pay.domain.models import (
    AcademicYear,
    HourlySelection,
    MonthlyAllocation,
    Teacher,
    WorkloadItem,
)
from hourly_pay.exporters.distribution_xlsx import ASSET, export_distribution


def _row(index: int, kind: WorkType) -> MonthlyPlanningRow:
    return MonthlyPlanningRow(
        selection=HourlySelection(index, Decimal("2"), id=index, teacher_id=1),
        item=WorkloadItem(
            id=index, study_form=StudyForm.FULL_TIME, source_row_number=index,
            group_name=f"Г-{index}", discipline=f"Дисциплина {index}",
            work_type=kind, hours=Decimal("2"), course="1",
        ),
        allocations=(MonthlyAllocation(index, 9, 2026, Decimal("2")),),
    )


def test_distribution_preserves_form_and_expands_all_work_types(tmp_path: Path) -> None:
    rows = tuple(
        [*(_row(index, WorkType.LECTURE) for index in range(1, 7)),
         *(_row(index, WorkType.PRACTICE) for index in range(7, 17)),
         _row(17, WorkType.LAB)]
    )
    funding = {index: "budget" if index % 2 else "commercial"
               for index in range(1, 18)}
    result = export_distribution(
        tmp_path, teacher=Teacher("Иванов Иван Иванович", department="Тестовая кафедра",
                                     position="доцент"),
        academic_year=AcademicYear("2026-2027"), rows=rows,
        funding_by_selection=funding,
    )
    assert result.row_count == 17
    assert result.hours_sum == Decimal("34")
    workbook = load_workbook(result.output_path)
    assert workbook.sheetnames == ["почасовка"]
    sheet = workbook.active
    assert sheet["B3"].value == "доцент"
    assert sheet["B5"].value == "Иванов"
    assert sheet["B6"].value == "Иван Иванович"
    assert sheet["P6"].value == 2026
    assert sheet["R6"].value == 2027
    assert sheet["Z5"].value == "Тестовая кафедра"
    assert [sheet.cell(row, 3).value for row in range(10, 30)
            if sheet.cell(row, 3).value] == [
                "Лекции", "Практические занятия", "Лабораторные занятия"
            ]
    assert sheet["E11"].value == "Дисциплина 1"
    assert sheet["G11"].value == 2
    assert sheet["H11"].value is None
    assert sheet["H12"].value == 2
    assert sheet["G12"].value is None
    assert sheet["AE11"].value.startswith("=SUM(")
    assert sheet["AF12"].value.startswith("=SUM(")
    total_row = 10 + 3 + 17
    assert sheet.cell(total_row, 5).value == "        Итого"
    assert sheet.cell(total_row, 7).value == f"=SUM(G11:G{total_row - 1})"
    assert sheet.print_area == f"'почасовка'!$A$1:$AG${total_row + 3}"


def test_distribution_requires_explicit_funding_and_complete_plan(tmp_path: Path) -> None:
    row = _row(1, WorkType.LAB)
    kwargs = dict(teacher=Teacher("Иванов И.И."),
                  academic_year=AcademicYear("2026-2027"), rows=(row,))
    with pytest.raises(InvalidDistributionExportError, match="категорию"):
        export_distribution(tmp_path, funding_by_selection={}, **kwargs)
    with pytest.raises(InvalidDistributionExportError, match="Помесячный план"):
        export_distribution(
            tmp_path, funding_by_selection={1: "budget"},
            **{**kwargs, "rows": (MonthlyPlanningRow(row.selection, row.item, ()),)},
        )
    wrong_year = MonthlyPlanningRow(
        row.selection, row.item,
        (MonthlyAllocation(1, 9, 2027, Decimal("2")),),
    )
    with pytest.raises(InvalidDistributionExportError, match="недопустимые месяцы"):
        export_distribution(
            tmp_path, funding_by_selection={1: "budget"},
            **{**kwargs, "rows": (wrong_year,)},
        )


def test_distribution_does_not_replace_existing_file(tmp_path: Path) -> None:
    kwargs = dict(teacher=Teacher("Иванов И.И."),
                  academic_year=AcademicYear("2026-2027"),
                  rows=(_row(1, WorkType.LECTURE),),
                  funding_by_selection={1: "budget"})
    first = export_distribution(tmp_path, **kwargs)
    original = first.output_path.read_bytes()
    with pytest.raises(InvalidDistributionExportError, match="уже существует"):
        export_distribution(tmp_path, **kwargs)
    assert first.output_path.read_bytes() == original


def test_packaged_distribution_template_is_sanitized() -> None:
    workbook = load_workbook(ASSET)
    assert workbook.sheetnames == ["почасовка"]
    assert workbook.active["B5"].value == "{FAMILIA}"
    assert workbook.active["Z5"].value == "{DEPARTMENT}"
