from datetime import date, time
from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document

from hourly_pay.application.schedule_month import ScheduleLessonRow, ScheduleMonth
from hourly_pay.application.statement_export import StatementExportService
from hourly_pay.domain.enums import WorkType
from hourly_pay.domain.errors import InvalidDocumentExportError
from hourly_pay.domain.models import Lesson, Teacher, TimeSlot
from hourly_pay.exporters.hourly_docx import DocxExportOptions


class ScheduleStub:
    def __init__(self, lesson_count: int) -> None:
        self.lesson_count = lesson_count
        self.validated = False
        slot = TimeSlot(1, 1, time(8, 30), time(10, 5), "08:30-10:05")
        lesson = Lesson(1, 1, date(2027, 3, 2), 1, "Г-1", "Дисциплина", WorkType.LECTURE, None)
        self.rows = (ScheduleLessonRow(lesson, slot),)
        self.slot = slot

    def available_months(self, teacher_id, academic_year_id):
        return (ScheduleMonth(2027, 3, self.lesson_count),)

    def validate_existing(self, teacher_id, academic_year_id, year, month):
        self.validated = True

    def list_lessons(self, teacher_id, year, month):
        return self.rows

    def list_time_slots(self):
        return (self.slot,)


class TeacherStub:
    def get_profile(self, teacher_id):
        from hourly_pay.application.teacher_settings import TeacherProfile

        return TeacherProfile(Teacher("Иванов Иван Иванович"), ())


def _template(path: Path) -> None:
    document = Document()
    document.add_paragraph("{FIO}")
    table = document.add_table(rows=2, cols=6)
    for cell, token in zip(
        table.rows[1].cells,
        ("{DATA}", "{TIME}", "{GROUP}", "{DISCIPLINE}", "{TYPE}", "{HOURS}"),
        strict=True,
    ):
        cell.text = token
    document.add_paragraph("Итого {HOURS_SUM}")
    document.save(path)


def test_statement_export_uses_saved_month_and_teacher(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _template(template)
    schedules = ScheduleStub(1)
    service = StatementExportService(schedules, TeacherStub())

    result = service.export(1, 2, 2027, 3, template, tmp_path / "out", DocxExportOptions())

    assert schedules.validated
    assert result.hours_sum == Decimal("2")
    assert result.lesson_count == 1
    assert Document(result.output_path).tables[0].rows[1].cells[3].text == "Дисциплина"


def test_statement_export_rejects_incomplete_month(tmp_path: Path) -> None:
    schedules = ScheduleStub(2)
    service = StatementExportService(schedules, TeacherStub())

    with pytest.raises(InvalidDocumentExportError, match="Количество занятий"):
        service.export(
            1, 2, 2027, 3, tmp_path / "template.docx", tmp_path / "out",
            DocxExportOptions(),
        )

    assert not (tmp_path / "out").exists()


def test_statement_export_rejects_month_outside_plan(tmp_path: Path) -> None:
    service = StatementExportService(ScheduleStub(1), TeacherStub())
    with pytest.raises(InvalidDocumentExportError, match="не относится"):
        service.export(
            1, 2, 2027, 4, tmp_path / "template.docx", tmp_path / "out",
            DocxExportOptions(),
        )
