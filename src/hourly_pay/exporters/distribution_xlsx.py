"""Fill the historical hourly-distribution form from saved monthly allocations."""

from collections.abc import Mapping, Sequence
from copy import copy
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from re import fullmatch, sub

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from hourly_pay.application.monthly_planning import MonthlyPlanningRow
from hourly_pay.domain.enums import WorkType
from hourly_pay.domain.errors import InvalidDistributionExportError
from hourly_pay.domain.models import AcademicYear, Teacher

MONTH_COLUMNS = {9: 7, 10: 9, 11: 11, 12: 13, 1: 15, 2: 17, 3: 19, 4: 21, 5: 23}
KINDS = (
    (WorkType.LECTURE, "Лекции"),
    (WorkType.PRACTICE, "Практические занятия"),
    (WorkType.LAB, "Лабораторные занятия"),
)
FUNDING_OFFSETS = {"budget": 0, "commercial": 1}
ASSET = Path(__file__).resolve().parents[1] / "assets" / "distribution_template.xlsx"


@dataclass(frozen=True, slots=True)
class DistributionExportResult:
    output_path: Path
    row_count: int
    hours_sum: Decimal


def export_distribution(
    output_directory: Path,
    *,
    teacher: Teacher,
    academic_year: AcademicYear,
    rows: Sequence[MonthlyPlanningRow],
    funding_by_selection: Mapping[int, str],
    template_path: Path = ASSET,
) -> DistributionExportResult:
    """Create one XLSX for the whole academic year, without guessing funding type."""
    match = fullmatch(r"(\d{4})[-–](\d{4})", academic_year.name.strip())
    if match is None or int(match.group(2)) != int(match.group(1)) + 1:
        raise InvalidDistributionExportError("Не распознан учебный год для распределения.")
    first_year, second_year = int(match.group(1)), int(match.group(2))
    if not rows:
        raise InvalidDistributionExportError("Сначала выберите нагрузку для почасовки.")
    selection_ids = {row.selection.id for row in rows}
    if None in selection_ids:
        raise InvalidDistributionExportError("Выбор нагрузки не сохранён.")
    if set(funding_by_selection) != selection_ids:
        raise InvalidDistributionExportError(
            "Для каждой строки выберите категорию: бюджет или коммерция."
        )
    if any(value not in FUNDING_OFFSETS for value in funding_by_selection.values()):
        raise InvalidDistributionExportError("Неизвестная категория финансирования.")

    hours_sum = Decimal("0")
    for row in rows:
        allocated = sum((entry.hours for entry in row.allocations), Decimal("0"))
        if allocated != row.selection.selected_hours:
            raise InvalidDistributionExportError(
                "Помесячный план не совпадает с выбранными часами. Сначала сохраните план."
            )
        if any(entry.month not in MONTH_COLUMNS or entry.hours <= 0
               or entry.hours % Decimal("2") != 0
               or entry.year != (first_year if entry.month >= 9 else second_year)
               for entry in row.allocations):
            raise InvalidDistributionExportError("В плане есть недопустимые месяцы или часы.")
        if len({entry.month for entry in row.allocations}) != len(row.allocations):
            raise InvalidDistributionExportError("Один месяц указан в плане несколько раз.")
        if row.item.work_type not in {kind for kind, _ in KINDS}:
            raise InvalidDistributionExportError("В выборе есть недопустимый вид занятия.")
        hours_sum += allocated

    if not template_path.is_file():
        raise InvalidDistributionExportError("Шаблон распределения XLSX не найден в приложении.")
    workbook = load_workbook(template_path)
    if workbook.sheetnames != ["почасовка"]:
        raise InvalidDistributionExportError("Формат шаблона распределения не распознан.")
    sheet = workbook.active
    if sheet["C10"].value != " Лекции" or sheet["C16"].value != "Практические занятия":
        raise InvalidDistributionExportError("Формат шаблона распределения не распознан.")

    section_styles = [copy(sheet.cell(16, col)._style) for col in range(1, 34)]
    data_styles = [copy(sheet.cell(17, col)._style) for col in range(1, 34)]
    sheet.unmerge_cells("C10:E10")
    sheet.unmerge_cells("C16:E16")
    sheet.delete_rows(10, 16)
    grouped = [(label, [row for row in rows if row.item.work_type is kind])
               for kind, label in KINDS]
    inserted = sum(1 + len(entries) for _, entries in grouped)
    sheet.insert_rows(10, inserted)

    name_parts = teacher.full_name.strip().split()
    sheet["B3"] = teacher.position or ""
    sheet["B5"] = name_parts[0] if name_parts else ""
    sheet["B6"] = " ".join(name_parts[1:])
    sheet["P6"] = int(match.group(1))
    sheet["R6"] = int(match.group(2))
    sheet["Z5"] = teacher.department or ""

    current_row = 10
    data_row_count = 0
    for label, entries in grouped:
        for col, style in enumerate(section_styles, start=1):
            sheet.cell(current_row, col)._style = copy(style)
        sheet.merge_cells(start_row=current_row, start_column=3,
                          end_row=current_row, end_column=5)
        sheet.cell(current_row, 3).value = label
        sheet.row_dimensions[current_row].height = 16
        current_row += 1
        for entry in entries:
            for col, style in enumerate(data_styles, start=1):
                sheet.cell(current_row, col)._style = copy(style)
            sheet.row_dimensions[current_row].height = 22
            sheet.cell(current_row, 4).value = entry.item.course or ""
            sheet.cell(current_row, 5).value = entry.item.discipline or ""
            sheet.cell(current_row, 6).value = entry.item.group_name or ""
            offset = FUNDING_OFFSETS[funding_by_selection[entry.selection.id]]
            for allocation in entry.allocations:
                column = MONTH_COLUMNS[allocation.month] + offset
                sheet.cell(current_row, column).value = int(allocation.hours)
            for total_col, month_offset in ((31, 0), (32, 1)):
                references = ",".join(
                    f"{get_column_letter(column + month_offset)}{current_row}"
                    for column in MONTH_COLUMNS.values()
                )
                sheet.cell(current_row, total_col).value = f"=SUM({references})"
            current_row += 1
            data_row_count += 1

    total_row = current_row
    first_data_row = 11
    for column in range(7, 33):
        letter = get_column_letter(column)
        sheet.cell(total_row, column).value = (
            f"=SUM({letter}{first_data_row}:{letter}{total_row - 1})"
        )
    sheet.row_dimensions[total_row].height = 16
    sheet.row_dimensions[total_row + 2].height = 18
    sheet.print_area = f"A1:AG{total_row + 3}"
    sheet.print_title_rows = "1:9"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    output_directory = Path(output_directory)
    output_directory.mkdir(parents=True, exist_ok=True)
    surname = sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name_parts[0]) if name_parts else "Преподаватель"
    filename = f"Распределение_почасовки_{surname}_{academic_year.name}.xlsx"
    output_path = output_directory / filename
    try:
        with output_path.open("xb") as stream:
            workbook.save(stream)
    except FileExistsError as error:
        raise InvalidDistributionExportError(
            f"Выходной файл уже существует: {filename}"
        ) from error
    except Exception:
        output_path.unlink(missing_ok=True)
        raise
    return DistributionExportResult(output_path, data_row_count, hours_sum)
