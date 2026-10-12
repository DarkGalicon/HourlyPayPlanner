from __future__ import annotations

import re
import shutil
from collections.abc import Mapping, Sequence
from copy import deepcopy
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from docx.document import Document as DocumentObject
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Row
from docx.shared import Pt
from docx.table import Table, _Cell, _Row
from docx.text.paragraph import Paragraph

from hourly_pay.config.defaults import DEFAULT_DOCX_WORK_TYPE_LABELS
from hourly_pay.domain.enums import WorkType
from hourly_pay.domain.errors import InvalidDocumentExportError
from hourly_pay.domain.models import Lesson, Teacher, TimeSlot

LESSON_PLACEHOLDERS = frozenset(
    {"{DATA}", "{TIME}", "{GROUP}", "{DISCIPLINE}", "{TYPE}", "{HOURS}"}
)
INVALID_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
MONTH_NAMES_GENITIVE = (
    "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
)


@dataclass(frozen=True, slots=True)
class DocxExportOptions:
    include_degree: bool = False
    include_title: bool = False
    include_position: bool = False
    statement_date: date | None = None
    work_type_labels: Mapping[WorkType, str] = field(
        default_factory=lambda: dict(DEFAULT_DOCX_WORK_TYPE_LABELS)
    )


@dataclass(frozen=True, slots=True)
class DocxExportResult:
    output_path: Path
    lesson_count: int
    hours_sum: Decimal


def export_hourly_statement(
    template_path: Path,
    output_directory: Path,
    *,
    teacher: Teacher,
    lessons: Sequence[Lesson],
    time_slots: Sequence[TimeSlot],
    year: int,
    month: int,
    options: DocxExportOptions | None = None,
) -> DocxExportResult:
    """Copy and fill a statement template without modifying the source DOCX."""
    template_path = Path(template_path)
    output_directory = Path(output_directory)
    options = options or DocxExportOptions()
    lesson_list = list(lessons)
    _validate_export(template_path, teacher, lesson_list, time_slots, year, month, options)

    output_directory.mkdir(parents=True, exist_ok=True)
    output_path = output_directory / safe_statement_filename(teacher.full_name, year, month)
    if output_path.resolve() == template_path.resolve():
        raise InvalidDocumentExportError("Исходный шаблон нельзя использовать как выходной файл.")
    if output_path.exists():
        raise InvalidDocumentExportError(f"Выходной файл уже существует: {output_path.name}")

    shutil.copy2(template_path, output_path)
    try:
        document = Document(output_path)
        _populate_lesson_table(document, lesson_list, time_slots, options.work_type_labels)
        _ensure_signature_table_spacing(document)
        hours_sum = sum((lesson.hours for lesson in lesson_list), start=Decimal("0"))
        _replace_document_placeholders(
            document,
            _document_values(teacher, hours_sum, options),
        )
        document.save(output_path)
    except (BadZipFile, PackageNotFoundError) as error:
        output_path.unlink(missing_ok=True)
        raise InvalidDocumentExportError("DOCX-шаблон повреждён или не распознан.") from error
    except Exception:
        output_path.unlink(missing_ok=True)
        raise

    return DocxExportResult(output_path, len(lesson_list), hours_sum)


def safe_statement_filename(full_name: str, year: int, month: int) -> str:
    parts = full_name.split()
    surname = parts[0] if parts else "Преподаватель"
    initials = "".join(_first_letter(part) for part in parts[1:3])
    name = "_".join(filter(None, (surname, initials)))
    name = INVALID_FILENAME_CHARS.sub("_", name).strip(" ._") or "Преподаватель"
    return f"Почасовка_{name}_{year:04d}-{month:02d}.docx"


def _first_letter(value: str) -> str:
    cleaned = "".join(character for character in value if character.isalpha())
    return cleaned[:1].upper()


def _validate_export(
    template_path: Path,
    teacher: Teacher,
    lessons: Sequence[Lesson],
    time_slots: Sequence[TimeSlot],
    year: int,
    month: int,
    options: DocxExportOptions,
) -> None:
    if not template_path.is_file() or template_path.suffix.lower() != ".docx":
        raise InvalidDocumentExportError("DOCX-шаблон не найден.")
    if not teacher.full_name.strip():
        raise InvalidDocumentExportError("Не указано ФИО преподавателя.")
    if not 1 <= month <= 12:
        raise InvalidDocumentExportError("Месяц должен быть от 1 до 12.")
    slot_ids = [slot.id for slot in time_slots]
    if len(slot_ids) != len(set(slot_ids)):
        raise InvalidDocumentExportError("В настройках времени есть повторяющиеся идентификаторы.")
    known_slots = set(slot_ids)
    for lesson in lessons:
        if lesson.lesson_date.year != year or lesson.lesson_date.month != month:
            raise InvalidDocumentExportError("Все занятия должны относиться к выбранному месяцу.")
        if lesson.hours != Decimal("2"):
            raise InvalidDocumentExportError(
                "Каждая строка заявления должна содержать ровно 2 часа."
            )
        if lesson.time_slot_id not in known_slots:
            raise InvalidDocumentExportError(
                f"Для занятия не найден временной слот: {lesson.time_slot_id}."
            )
        if lesson.work_type not in options.work_type_labels:
            raise InvalidDocumentExportError(
                f"Для вида занятия {lesson.work_type.value} не настроен текст DOCX."
            )


def _populate_lesson_table(
    document: DocumentObject,
    lessons: Sequence[Lesson],
    time_slots: Sequence[TimeSlot],
    labels: Mapping[WorkType, str],
) -> None:
    located = _find_lesson_row(document)
    if located is None:
        raise InvalidDocumentExportError("В DOCX-шаблоне не найдена строка таблицы занятий.")
    table, row_index = located
    prototype = table.rows[row_index]
    prototype_xml = deepcopy(prototype._tr)
    slot_text = {slot.id: slot.display_text for slot in time_slots}

    if row_index > 0:
        _mark_as_repeating_header(table.rows[row_index - 1])
    if not lessons:
        table._tbl.remove(prototype._tr)
        return

    for index, lesson in enumerate(lessons):
        row = (
            prototype
            if index == 0
            else _clone_row_after(table, table.rows[row_index + index - 1], prototype_xml)
        )
        _mark_row_as_non_splitting(row)
        values = {
            "{DATA}": lesson.lesson_date.strftime("%d.%m.%Y"),
            "{TIME}": slot_text[lesson.time_slot_id],
            "{GROUP}": lesson.group_name,
            "{DISCIPLINE}": lesson.discipline,
            "{TYPE}": labels[lesson.work_type],
            "{HOURS}": _format_decimal(lesson.hours),
        }
        _replace_in_row(row, values)
        if len(lessons) >= 7:
            for cell in row.cells:
                for paragraph in _cell_paragraphs(cell):
                    for run in paragraph.runs:
                        run.font.size = Pt(10)


def _find_lesson_row(document: DocumentObject) -> tuple[Table, int] | None:
    for table in document.tables:
        for index, row in enumerate(table.rows):
            row_text = "\n".join(cell.text for cell in row.cells)
            if LESSON_PLACEHOLDERS.issubset(
                {placeholder for placeholder in LESSON_PLACEHOLDERS if placeholder in row_text}
            ):
                return table, index
    return None


def _clone_row_after(table: Table, source: _Row, prototype_xml: CT_Row) -> _Row:
    cloned = deepcopy(prototype_xml)
    for paragraph in cloned.findall(".//" + qn("w:p")):
        paragraph.attrib.pop(qn("w14:paraId"), None)
        paragraph.attrib.pop(qn("w14:textId"), None)
    source._tr.addnext(cloned)
    return _Row(cloned, table)


def _mark_as_repeating_header(row: _Row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:tblHeader")) is None:
        properties.append(OxmlElement("w:tblHeader"))


def _mark_row_as_non_splitting(row: _Row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def _ensure_signature_table_spacing(document: DocumentObject) -> None:
    """Keep two empty paragraphs between the date/signature and hours tables."""
    for date_table, hours_table in zip(document.tables, document.tables[1:], strict=False):
        date_text = " ".join(cell.text for row in date_table.rows for cell in row.cells)
        hours_text = " ".join(cell.text for row in hours_table.rows for cell in row.cells)
        if "{MONTH}" not in date_text or "{YEAR}" not in date_text:
            continue
        if "Количество часов" not in hours_text:
            continue
        between = []
        element = date_table._tbl.getnext()
        while element is not None and element is not hours_table._tbl:
            if element.tag != qn("w:p") or Paragraph(element, document).text.strip():
                break
            between.append(element)
            element = element.getnext()
        if element is not hours_table._tbl:
            continue
        for paragraph in between:
            paragraph.getparent().remove(paragraph)
        for _ in range(2):
            element = OxmlElement("w:p")
            hours_table._tbl.addprevious(element)
            paragraph = Paragraph(element, document)
            paragraph.style = document.styles["Normal"]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(10)
        return


def _replace_in_row(row: _Row, values: Mapping[str, str]) -> None:
    for cell in row.cells:
        for paragraph in _cell_paragraphs(cell):
            _replace_in_paragraph(paragraph, values)


def _document_values(
    teacher: Teacher,
    hours_sum: Decimal,
    options: DocxExportOptions,
) -> dict[str, str]:
    statement_date = options.statement_date
    values = {
        "{FIO}": teacher.full_name,
        "{STEPEN}": teacher.degree or "" if options.include_degree else "",
        "{ZVANIE}": teacher.title or "" if options.include_title else "",
        "{DOLZNOST}": teacher.position or "" if options.include_position else "",
        "{BIRTH_DATE}": teacher.birth_date.strftime("%d.%m.%Y") if teacher.birth_date else "",
        "{ADDRESS}": teacher.address or "",
        "{PASS_SERIAL}": teacher.passport_series or "",
        "{PASS_NUM}": teacher.passport_number or "",
        "{OTDEL}": teacher.passport_issued_by or "",
        "{DATE_OTDEL}": (
            teacher.passport_issue_date.strftime("%d.%m.%Y")
            if teacher.passport_issue_date else ""
        ),
        "{STRAX_NUM}": teacher.snils or "",
        "{INN_NUM}": teacher.inn or "",
        "{HOURS_SUM}": _format_decimal(hours_sum),
        "{DAY)": f"{statement_date.day:02d}" if statement_date else "",
        "{DAY}": f"{statement_date.day:02d}" if statement_date else "",
        "{MONTH}": MONTH_NAMES_GENITIVE[statement_date.month - 1] if statement_date else "",
        "{YEAR}": f"{statement_date.year:04d}" if statement_date else "",
    }
    return values


def _replace_document_placeholders(
    document: DocumentObject, values: Mapping[str, str]
) -> None:
    for paragraph in document.paragraphs:
        _replace_in_paragraph(paragraph, values)
    for table in document.tables:
        for row in table.rows:
            _replace_in_row(row, values)
    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                _replace_in_paragraph(paragraph, values)
            for table in container.tables:
                for row in table.rows:
                    _replace_in_row(row, values)


def _cell_paragraphs(cell: _Cell) -> list[Paragraph]:
    paragraphs = list(cell.paragraphs)
    for table in cell.tables:
        for row in table.rows:
            for nested_cell in row.cells:
                paragraphs.extend(_cell_paragraphs(nested_cell))
    return paragraphs


def _replace_in_paragraph(paragraph: Paragraph, values: Mapping[str, str]) -> None:
    had_hours_sum = "{HOURS_SUM}" in "".join(run.text for run in paragraph.runs)
    for placeholder, value in values.items():
        _replace_token_in_runs(paragraph, placeholder, value)
    if had_hours_sum:
        _remove_word_field_markup(paragraph)


def _remove_word_field_markup(paragraph: Paragraph) -> None:
    for run_element in list(paragraph._p.findall(qn("w:r"))):
        if run_element.find(qn("w:fldChar")) is not None or run_element.find(
            qn("w:instrText")
        ) is not None:
            paragraph._p.remove(run_element)


def _replace_token_in_runs(paragraph: Paragraph, token: str, value: str) -> None:
    while True:
        runs = paragraph.runs
        complete_text = "".join(run.text for run in runs)
        start = complete_text.find(token)
        if start < 0:
            return
        end = start + len(token)
        cursor = 0
        first_index = last_index = -1
        first_offset = last_offset = 0
        for index, run in enumerate(runs):
            run_end = cursor + len(run.text)
            if first_index < 0 and start < run_end:
                first_index = index
                first_offset = start - cursor
            if end <= run_end:
                last_index = index
                last_offset = end - cursor
                break
            cursor = run_end
        if first_index < 0 or last_index < 0:
            return
        first_run = runs[first_index]
        last_run = runs[last_index]
        prefix = first_run.text[:first_offset]
        suffix = last_run.text[last_offset:]
        if first_index == last_index:
            first_run.text = f"{prefix}{value}{suffix}"
            continue
        first_run.text = f"{prefix}{value}"
        for run in runs[first_index + 1 : last_index]:
            run.text = ""
        last_run.text = suffix


def _format_decimal(value: Decimal) -> str:
    rendered = format(value, "f")
    return rendered.rstrip("0").rstrip(".") if "." in rendered else rendered
