from __future__ import annotations

import logging
import sys
from collections.abc import Sequence
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication

from hourly_pay.application.availability_settings import AvailabilitySettingsService
from hourly_pay.application.calculate_teacher import TeacherCalculationService
from hourly_pay.application.calendar_settings import CalendarSettingsService
from hourly_pay.application.import_workload import ImportWorkloadService
from hourly_pay.application.monthly_planning import MonthlyPlanningService
from hourly_pay.application.schedule_month import ScheduleMonthService
from hourly_pay.application.selection_service import SelectionService
from hourly_pay.application.semester_settings import SemesterSettingsService
from hourly_pay.application.statement_export import StatementExportService
from hourly_pay.application.teacher_settings import TeacherSettingsService
from hourly_pay.exporters.distribution_xlsx import ASSET as DISTRIBUTION_TEMPLATE
from hourly_pay.persistence.db import create_session_factory, create_sqlite_engine, upgrade_database
from hourly_pay.persistence.repositories import (
    CalendarOverrideRepository,
    HourlySelectionRepository,
    MonthlyAllocationRepository,
    ScheduleRepository,
    SemesterStartRepository,
    WorkloadRepository,
)
from hourly_pay.ui.main_window import MainWindow
from hourly_pay.ui.theme import apply_application_appearance


def create_main_window(database_path: Path | None = None) -> MainWindow:
    resolved_path = upgrade_database(database_path)
    engine = create_sqlite_engine(resolved_path)
    repository = WorkloadRepository(create_session_factory(engine))
    calculation_service = TeacherCalculationService(repository)
    selection_repository = HourlySelectionRepository(create_session_factory(engine))
    allocation_repository = MonthlyAllocationRepository(create_session_factory(engine))
    schedule_repository = ScheduleRepository(create_session_factory(engine))
    calendar_repository = CalendarOverrideRepository(create_session_factory(engine))
    calendar_service = CalendarSettingsService(calendar_repository)
    semester_service = SemesterSettingsService(
        SemesterStartRepository(create_session_factory(engine))
    )
    monthly_service = MonthlyPlanningService(
        calculation_service, selection_repository, allocation_repository
    )
    teacher_settings = TeacherSettingsService(repository)
    schedule_service = ScheduleMonthService(
        monthly_service, schedule_repository, calendar_service, semester_service
    )
    window = MainWindow(
        ImportWorkloadService(repository),
        teacher_settings,
        calculation_service,
        SelectionService(calculation_service, selection_repository),
        monthly_service,
        AvailabilitySettingsService(repository, schedule_repository),
        calendar_service,
        schedule_service,
        semester_service,
        StatementExportService(schedule_service, teacher_settings),
    )
    window.database_engine = engine
    return window


def smoke_test() -> int:
    """Exercise packaged Qt, migrations, and DOCX resources without touching user data."""
    application = QApplication(["HourlyPayPlanner", "-platform", "offscreen"])
    application.setApplicationName("Hourly Pay Planner")
    apply_application_appearance(application)
    with TemporaryDirectory(prefix="hourly-pay-smoke-") as directory:
        window = create_main_window(Path(directory) / "smoke.db")
        try:
            distribution_template = load_workbook(DISTRIBUTION_TEMPLATE, read_only=True)
            try:
                template_ready = distribution_template.sheetnames == ["почасовка"]
            finally:
                distribution_template.close()
            ready = (
                window.documents_page is not None
                and window.schedule_page is not None
                and template_ready
            )
            document = Document()
            document.sections[0].header.paragraphs[0].text = "Проверка шапки"
            document.sections[0].footer.paragraphs[0].text = "Проверка подвала"
            document.save(Path(directory) / "smoke.docx")
        finally:
            window.close()
            window.database_engine.dispose()
    application.quit()
    return 0 if ready else 1


def main(argv: Sequence[str] | None = None) -> int:
    logging.basicConfig(level=logging.INFO)
    application = QApplication(list(argv) if argv is not None else sys.argv)
    application.setApplicationName("Hourly Pay Planner")
    application.setOrganizationName("HourlyPayPlanner")
    apply_application_appearance(application)
    window = create_main_window()
    window.show()
    return application.exec()
